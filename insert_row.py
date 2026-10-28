#!/usr/bin/env python3
"""Insert a property row into a month block, creating the month header if needed.

Rebuilds the sheet as a permutation of its rows plus the new ones, so Rate PM2
formulas and merged month headers are remapped rather than left dangling -- the
same approach move_rows.py and update_workbook.py use.

    python3 insert_row.py <repo dir> <workbook>
"""
import sys, copy, re, json
import openpyxl
from openpyxl.styles import Font
sys.path.insert(0, sys.argv[1])
import update_workbook as U

WB = sys.argv[2]

# sheet, month index, [addr, price, bed, bath, car, size]
INSERTS = [
    ('Point Piper UNIT', 10, ['7/6 Wentworth Street', None, '-', '-', '-', 192]),
]


def write_prop(ws, row, tmpl_cells, vals, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell._style = copy.copy(tmpl_cells[c]['style'])
        cell.value = vals[c - 1] if c <= len(vals) else None
        cell.number_format = U.MONEY if c in (2, 7) else tmpl_cells[c]['fmt']
        cell.hyperlink = None
        cell.font = Font(name='Arial', size=19,
                         color=(U.BLUE if c == 1 else 'FF000000'))


def write_hdr(ws, row, tmpl_cells, text, max_col):
    for c in range(1, max_col + 1):
        cell = ws.cell(row=row, column=c)
        cell._style = copy.copy(tmpl_cells[c]['style'])
        cell.value = text if c == 1 else None
        cell.number_format = tmpl_cells[c]['fmt']
        cell.hyperlink = None


wb = openpyxl.load_workbook(WB)
log = []
for sheet, midx, vals in INSERTS:
    ws = wb[sheet]
    sh = U.Sheet(ws)
    max_col = sh.max_col
    snap = U.snapshot(ws, sh.max_row, max_col)
    blk = next((b for b in sh.blocks if b[1] == midx), None)

    hdr_txt = None
    created = None
    if blk:                                   # block exists: append into it
        anchor, marks = blk[4], ['NEW']
        tmpl_prop = sh.tmpl_by_block.get(blk[0]) or sh.tmpl_prop
        if blk[3] is not None:
            hdr_txt = re.sub(r'\(\d+\)', '(%d)' % (blk[3] + 1),
                             str(snap[blk[0]]['cells'][1]['v']))
    else:                                     # create the header, in order
        later = [b for b in sh.blocks if b[1] is not None and b[1] > midx]
        anchor = (later[0][0] - 1) if later else sh.last_data
        marks = ['HDR', 'NEW']
        tmpl_prop = sh.tmpl_prop
        created = '%s (1)' % U.MONTHS[midx - 1]

    order = []
    for r in range(1, sh.max_row + 1):
        order.append(r)
        if r == anchor:
            order.extend(marks)
    omap = {old: new for new, old in enumerate(order, start=1)
            if not isinstance(old, str)}

    old_merges = [(m.min_row, m.min_col, m.max_row, m.max_col)
                  for m in list(ws.merged_cells.ranges)]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))

    new_hdr_row = None
    for new_r, old_r in enumerate(order, start=1):
        if old_r == 'NEW':
            write_prop(ws, new_r, snap[tmpl_prop]['cells'], vals, max_col)
            new_prop_row = new_r
            continue
        if old_r == 'HDR':
            write_hdr(ws, new_r, snap[sh.tmpl_month]['cells'], created, max_col)
            new_hdr_row = new_r
            continue
        s = snap[old_r]
        for c in range(1, max_col + 1):
            cell = ws.cell(row=new_r, column=c)
            cd = s['cells'][c]
            v = cd['v']
            if isinstance(v, str) and v.startswith('='):
                v = U.remap_formula(v, omap)
            if c == 1 and blk and old_r == blk[0] and hdr_txt:
                v = hdr_txt
            cell.value = v
            cell._style = copy.copy(cd['style'])
            cell.number_format = cd['fmt']
            cell.hyperlink = copy.copy(cd['link']) if cd['link'] else None
        if s['h'] is not None:
            ws.row_dimensions[new_r].height = s['h']

    for mnr, mnc, mxr, mxc in old_merges:
        if mnr == mxr and mnr in omap:
            ws.merge_cells(start_row=omap[mnr], start_column=mnc,
                           end_row=omap[mnr], end_column=mxc)
    if new_hdr_row:                            # a month header spans A:G
        ws.merge_cells(start_row=new_hdr_row, start_column=1,
                       end_row=new_hdr_row, end_column=7)

    log.append({'sheet': sheet, 'addr': vals[0], 'month': U.MONTHS[midx - 1],
                'new_row': new_prop_row, 'header_created': created,
                'header_row': new_hdr_row, 'header_updated': hdr_txt,
                'styled_after_row': tmpl_prop})

wb.save(WB)
print(json.dumps(log, indent=1))
