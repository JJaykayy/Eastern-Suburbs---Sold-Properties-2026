#!/usr/bin/env python3
"""Weekly RP Data export -> "Eastern Suburbs - Sold Properties 2026.xlsx" updater.

Usage, from the folder holding the workbook:
    python3 update_workbook.py <export file> --apply [--type-map map.json] [--report r.json]

Omit --apply for a dry run that only writes the report. Reads .numbers, .csv,
.tsv, .xlsx exports, with or without a header row.

Requires: openpyxl, and numbers-parser for .numbers exports.
"""
import re, sys, json, copy, difflib
from collections import defaultdict, OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

WB_PATH = 'Eastern Suburbs - Sold Properties 2026.xlsx'
OFF_LIMITS = {'suburb details', 'vacancy'}
MONTHS = ['January','February','March','April','May','June','July','August',
          'September','October','November','December']
MIDX = {m.upper(): i+1 for i, m in enumerate(MONTHS)}
MABBR = {m[:3].upper(): i+1 for i, m in enumerate(MONTHS)}
MONTH_RE = re.compile(r'^\s*(' + '|'.join(MONTHS) + r')\s*\((\d+)\)\s*$', re.I)
# Some month headers in this workbook carry no "(count)" at all. They are still
# month headers (bold, 19pt, #93C47D, merged A:G) and must never be duplicated.
GREEN = 'FF93C47D'          # month-header fill
BLUE = 'FF4285F4'           # address font colour
MONEY = '"$"#,##0'          # Sale Price / Rate PM2 number format


def is_green(cell):
    fill = cell.fill
    return (fill is not None and fill.patternType == 'solid'
            and fill.fgColor is not None and fill.fgColor.rgb == GREEN)


def month_header(cell, txt):
    """-> (month_idx|None, count|None) if this row is a month header, else None.

    A month header is recognised by the green header fill, so headers that lack a
    "(count)" or whose month name is misspelled ('Januray (1)') are still headers
    and never get a duplicate created alongside them."""
    m = MONTH_RE.match(txt)
    if m:
        return MIDX[m.group(1).upper()], int(m.group(2))
    if 'vacancy' in txt.lower():
        return None
    if not is_green(cell):
        return None
    cnt = None
    mc = re.search(r'\((\d+)\)', txt)
    if mc:
        cnt = int(mc.group(1))
    word = re.sub(r'[^A-Za-z]', ' ', txt).split()
    midx = None
    if word:
        close = difflib.get_close_matches(word[0].capitalize(), MONTHS, n=1, cutoff=0.6)
        if close:
            midx = MIDX[close[0].upper()]
    return midx, cnt

SUF = {'ST':'STREET','STR':'STREET','RD':'ROAD','AVE':'AVENUE','AV':'AVENUE','AVEN':'AVENUE',
       'CRES':'CRESCENT','CRESENT':'CRESCENT','CR':'CRESCENT','CRE':'CRESCENT','PDE':'PARADE',
       'DR':'DRIVE','DRV':'DRIVE','DVE':'DRIVE','LN':'LANE','PL':'PLACE','CT':'COURT',
       'TCE':'TERRACE','TER':'TERRACE','TRCE':'TERRACE','BLVD':'BOULEVARD','BVD':'BOULEVARD',
       'BOULEVARDE':'BOULEVARD','HWY':'HIGHWAY','SQ':'SQUARE','GDN':'GARDENS','GDNS':'GARDENS',
       'CL':'CLOSE','WY':'WAY','GR':'GROVE','GRV':'GROVE','ESP':'ESPLANADE','CCT':'CIRCUIT',
       'CIRCT':'CIRCUIT','PKWY':'PARKWAY','PWY':'PARKWAY','RW':'ROW','CIR':'CIRCLE'}
DESCRIPTORS = {'UNIT','APT','APARTMENT','VILLA','TERRACE','TOWNHOUSE','SHOP','SUITE','LOT',
               'STUDIO','PENTHOUSE','LEVEL','HOUSE','DUPLEX','FLAT'}
NUM = r'\d+[A-Z]?'
APAT = re.compile(r'^(?:([A-Z]{0,2}\d+[A-Z]?)\s*/\s*)?(' + NUM + r'(?:\s*-\s*' + NUM + r')?)\s+(.+)$')


def norm_text(s):
    s = str(s).upper().replace(' ', ' ')
    s = s.replace(',', ' ').replace('.', ' ').replace('  ', ' ')
    s = re.sub(r'\s+', ' ', s).strip()
    toks = s.split()
    if toks and toks[-1] in SUF:
        toks[-1] = SUF[toks[-1]]
    return ' '.join(toks)


def norm_unit(u):
    if u is None:
        return None
    m = re.match(r'^([A-Z]*)0*(\d+)([A-Z]?)$', u)
    if not m:
        return u
    return (m.group(1) or '') + m.group(2) + (m.group(3) or '')


def num_range(numtxt):
    parts = [p.strip() for p in numtxt.split('-')]
    vals = []
    for p in parts:
        m = re.match(r'^(\d+)([A-Z]?)$', p)
        vals.append((int(m.group(1)), m.group(2)))
    if len(vals) == 1:
        return vals[0][0], vals[0][0], vals[0][1], True
    lo = min(v[0] for v in vals); hi = max(v[0] for v in vals)
    return lo, hi, '', False


def parse_addr(raw):
    n = norm_text(raw)
    toks = n.split()
    while len(toks) > 2 and toks[0] in DESCRIPTORS and APAT.match(' '.join(toks[1:])):
        toks = toks[1:]
    rest = ' '.join(toks)
    m = APAT.match(rest)
    if not m:
        return {'full': n, 'street': None}
    unit, numtxt, street = m.groups()
    lo, hi, sfx, single = num_range(numtxt)
    return {'full': n, 'unit': norm_unit(unit), 'street': street.strip(),
            'lo': lo, 'hi': hi, 'sfx': sfx, 'single': single}


def same_property(a, b):
    if a['full'] == b['full']:
        return True
    if a.get('street') is None or b.get('street') is None:
        return False
    if a['unit'] != b['unit'] or a['street'] != b['street']:
        return False
    if not (a['lo'] <= b['hi'] and b['lo'] <= a['hi']):
        return False
    if a['single'] and b['single'] and a['sfx'] != b['sfx']:
        return False
    return True


def title_addr(s):
    """Match the workbook's Title Case addresses (the export is all caps).
    Tokens containing digits keep their exact source form ('36B', '5B/356-368')."""
    def fix(tok):
        if any(ch.isdigit() for ch in tok):
            return tok
        t = re.sub(r'[A-Za-z]+',
                   lambda m: m.group(0)[0].upper() + m.group(0)[1:].lower(), tok)
        m = re.match(r'^(Ma?c)([a-z])(.*)$', t)
        if m and len(t) > 3:
            t = m.group(1) + m.group(2).upper() + m.group(3)
        return t
    return ' '.join(fix(t) for t in str(s).split())


def clean_num(v):
    if v is None:
        return None
    s = str(v).strip()
    if s in ('', '-', 'Not Disclosed', 'N/A', 'NA'):
        return None
    s = s.replace('$', '').replace(',', '')
    try:
        f = float(s)
    except ValueError:
        return None
    if f <= 0:
        return None
    return int(round(f)) if abs(f - round(f)) < 1e-9 else f


def clean_price(v):
    """A usable sale price. $0 / $1 are RP Data placeholders, not prices: leaving
    the cell blank also keeps the row eligible for a real price in a later export."""
    n = clean_num(v)
    if n is None or n <= 1:
        return None
    return n


def small_num(v):
    """bed/bath/car: number, explicit '-' for none, or None for unknown."""
    if v is None:
        return None
    s = str(v).strip()
    if s == '':
        return None
    if s == '-':
        return '-'
    try:
        f = float(s)
    except ValueError:
        return None
    return int(round(f)) if abs(f - round(f)) < 1e-9 else f


def parse_date(v):
    s = str(v).strip()
    m = re.match(r'^(\d{1,2})\s+([A-Za-z]{3,})\s+(\d{4})$', s)
    if m:
        mi = MABBR.get(m.group(2)[:3].upper())
        if mi:
            return int(m.group(3)), mi, int(m.group(1))
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if m:
        return int(m.group(1)), int(m.group(2)), int(m.group(3))
    m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})$', s)
    if m:
        return int(m.group(3)), int(m.group(2)), int(m.group(1))
    return None


def sibling_key(target):
    """'ZETLAND HOUSE' <-> 'ZETLAND UNIT' - the other sheet for the same suburb."""
    if target.endswith(' HOUSE'):
        return target[:-6] + ' UNIT'
    if target.endswith(' UNIT'):
        return target[:-5] + ' HOUSE'
    return None


def sheet_key(name):
    n = re.sub(r'[^A-Z ]', ' ', str(name).upper())
    n = re.sub(r'\s+', ' ', n).strip()
    return n


# ---------------------------------------------------------------- read export
def read_export(path):
    if path.lower().endswith('.numbers'):
        from numbers_parser import Document
        doc = Document(path)
        rows = doc.sheets[0].tables[0].rows(values_only=True)
        rows = [list(r) for r in rows]
    elif path.lower().endswith(('.csv', '.tsv', '.txt')):
        import csv
        delim = '\t' if path.lower().endswith('.tsv') else ','
        with open(path, newline='', encoding='utf-8-sig') as fh:
            rows = [r for r in csv.reader(fh, delimiter=delim)]
    elif path.lower().endswith(('.xlsx', '.xlsm')):
        wb = openpyxl.load_workbook(path, data_only=True)
        rows = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    else:
        raise SystemExit('unsupported export format: ' + path)
    rows = [r for r in rows if any(c not in (None, '') for c in r)]
    # header detection
    COLS = ['address','suburb','state','postcode','property type','bed','bath','car',
            'land size','internal','year built','sale price','sale date','settlement',
            'sale type','agency']
    idx = list(range(16))
    if rows:
        first = [str(c).strip().lower() if c is not None else '' for c in rows[0]]
        hits = sum(1 for c in first if any(k in c for k in
                   ('address','suburb','price','bed','bath','agency','sale date')))
        if hits >= 3:
            def find(*keys):
                for i, c in enumerate(first):
                    if any(k in c for k in keys):
                        return i
                return None
            cand = [find('address'), find('suburb'), find('state'), find('postcode'),
                    find('property type','prop type','type'), find('bed'), find('bath'),
                    find('car','park'), find('land'), find('internal','floor','building'),
                    find('year'), find('sale price','price'), find('sale date','sold date'),
                    find('settlement'), find('sale type','status'), find('agency','agent')]
            if cand[0] is not None and cand[11] is not None:
                idx = [c if c is not None else -1 for c in cand]
            rows = rows[1:]
    out = []
    for r in rows:
        r = list(r) + [None] * 16
        out.append([(r[i] if 0 <= i < len(r) else None) for i in idx])
    return out


# ------------------------------------------------------------ workbook model
class Sheet:
    def __init__(self, ws):
        self.ws = ws
        self.donor = None        # donor sheet name for a freshly created sheet
        self.max_row = ws.max_row
        self.max_col = max(7, ws.max_column)
        self.last_data = 1
        for r in range(1, self.max_row + 1):
            if any(ws.cell(row=r, column=c).value not in (None, '')
                   for c in range(1, self.max_col + 1)):
                self.last_data = r
        self.kind = {}
        self.props = []          # (row, parsed_addr)
        self.headers = []        # (row, month_idx, month_name, count)
        for r in range(1, self.last_data + 1):
            a = self.ws.cell(row=r, column=1).value
            txt = '' if a is None else str(a)
            if r == 1:
                self.kind[r] = 'title'
                continue
            mh = month_header(self.ws.cell(row=r, column=1), txt)
            if mh:
                self.kind[r] = 'month'
                self.headers.append([r, mh[0],
                                     MONTHS[mh[0]-1] if mh[0] else txt.strip(), mh[1]])
            elif 'vacancy' in txt.lower():
                self.kind[r] = 'vacancy'
            elif txt.strip() == '':
                self.kind[r] = 'blank'
            else:
                self.kind[r] = 'prop'
                self.props.append((r, parse_addr(txt)))
        self.blocks = []  # [hdr_row, midx, name, count, end_row]
        for i, h in enumerate(self.headers):
            end = (self.headers[i+1][0] - 1) if i + 1 < len(self.headers) else self.last_data
            self.blocks.append(h + [end])
        # style template: a property row that is not carrying header fill
        self.tmpl_prop = None
        for r, _ in self.props:
            if not is_green(self.ws.cell(row=r, column=1)):
                self.tmpl_prop = r
                break
        if self.tmpl_prop is None and self.props:
            self.tmpl_prop = self.props[0][0]
        self.tmpl_by_block = {}
        for b in self.blocks:
            cand = [r for r, _ in self.props if b[0] < r <= b[4]
                    and not is_green(self.ws.cell(row=r, column=1))]
            if cand:
                self.tmpl_by_block[b[0]] = cand[-1]
        self.tmpl_month = None
        for h in self.headers:
            if h[1] is not None and h[3] is not None:
                self.tmpl_month = h[0]
                break
        if self.tmpl_month is None and self.headers:
            self.tmpl_month = self.headers[0][0]

    def find(self, parsed):
        for r, p in self.props:
            if same_property(parsed, p):
                return r
        return None


CELLREF = re.compile(r'(\$?)([A-Z]{1,3})(\$?)(\d+)')


def remap_formula(f, omap):
    def rep(m):
        r = int(m.group(4))
        return m.group(1) + m.group(2) + m.group(3) + str(omap.get(r, r))
    return CELLREF.sub(rep, f)


def snapshot(ws, max_row, max_col):
    snap = {}
    for r in range(1, max_row + 1):
        cells = {}
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cells[c] = {'v': cell.value, 'style': copy.copy(cell._style),
                        'fmt': cell.number_format,
                        'link': copy.copy(cell.hyperlink) if cell.hyperlink else None}
        snap[r] = {'cells': cells, 'h': ws.row_dimensions[r].height
                   if r in ws.row_dimensions else None}
    return snap


def write_changes(wb, model, plans, path, report):
    report['verify'] = []
    for name, plan in plans.items():
        if not plan['new'] and not plan['fills']:
            continue
        sh = model[name]
        ws = sh.ws
        max_col = sh.max_col
        snap = snapshot(ws, sh.max_row, max_col)

        # ---- insertion planning -------------------------------------------
        by_month = defaultdict(list)
        for rec in plan['new']:
            by_month[rec['date'][1]].append(rec)
        existing = {b[1]: b for b in sh.blocks if b[1] is not None}
        inserts = defaultdict(list)
        hdr_updates = {}
        created_months = []
        countless = []
        for midx in sorted(by_month):
            recs = by_month[midx]
            if midx in existing:
                blk = existing[midx]
                anchor = blk[4] + 1
                for rec in recs:
                    rec['_tmpl'] = sh.tmpl_by_block.get(blk[0]) or sh.tmpl_prop
                if blk[3] is None:
                    # pre-existing header with no "(count)": append into the block
                    # but leave its text alone rather than inventing a count.
                    countless.append(f'{MONTHS[midx-1]} (no count in sheet)')
                else:
                    hdr_updates[blk[0]] = blk[3] + len(recs)
                for rec in recs:
                    inserts[anchor].append(('prop', rec))
            else:
                later = [b for b in sh.blocks if b[1] is not None and b[1] > midx]
                anchor = later[0][0] if later else sh.last_data + 1
                inserts[anchor].append(('month', {'midx': midx, 'count': len(recs)}))
                for rec in recs:
                    inserts[anchor].append(('prop', rec))
                created_months.append(MONTHS[midx - 1])

        target = []
        for r in range(1, sh.max_row + 1):
            for it in inserts.get(r, []):
                target.append(it)
            target.append(('old', r))
        for it in inserts.get(sh.max_row + 1, []):
            target.append(it)

        omap = {it[1]: i for i, it in enumerate(target, start=1) if it[0] == 'old'}
        fills = {f['row']: f for f in plan['fills']}

        # templates (snapshot rows, so they are pre-shift). A sheet created this
        # run has only a header row, so its templates come from a donor sheet.
        tprop, tmonth = sh.tmpl_prop, sh.tmpl_month
        tsnap = snap
        if (tprop is None or tmonth is None) and getattr(sh, 'donor', None):
            dws = wb[sh.donor]
            dsh = Sheet(dws)
            tsnap = snapshot(dws, dsh.max_row, max(max_col, dsh.max_col))
            tprop = tprop or dsh.tmpl_prop
            tmonth = tmonth or dsh.tmpl_month

        # ---- unmerge everything before writing ---------------------------
        old_merges = [(m.min_row, m.min_col, m.max_row, m.max_col)
                      for m in list(ws.merged_cells.ranges)]
        for m in list(ws.merged_cells.ranges):
            ws.unmerge_cells(str(m))

        new_month_rows = []
        touched_hdr_rows = []
        added_rows = []

        for new_r, item in enumerate(target, start=1):
            if item[0] == 'old':
                old_r = item[1]
                src = snap[old_r]
                for c in range(1, max_col + 1):
                    sc = src['cells'][c]
                    cell = ws.cell(row=new_r, column=c)
                    v = sc['v']
                    if isinstance(v, str) and v.startswith('='):
                        v = remap_formula(v, omap)
                    cell.value = v
                    cell._style = copy.copy(sc['style'])
                    cell.number_format = sc['fmt']
                    if sc['link'] is not None:
                        h = sc['link']
                        cell.hyperlink = Hyperlink(
                            ref=f'{get_column_letter(c)}{new_r}', target=h.target,
                            tooltip=h.tooltip, display=h.display,
                            location=h.location)
                    else:
                        cell.hyperlink = None
                if src['h'] is not None:
                    ws.row_dimensions[new_r].height = src['h']
                if old_r in hdr_updates:
                    blk = next(b for b in sh.blocks if b[0] == old_r)
                    ws.cell(row=new_r, column=1).value = \
                        f'{MONTHS[blk[1]-1]} ({hdr_updates[old_r]})'
                    touched_hdr_rows.append((new_r, MONTHS[blk[1]-1], hdr_updates[old_r]))
                if old_r in fills:
                    f = fills[old_r]
                    b = ws.cell(row=new_r, column=2)
                    b.value = f['price']
                    b.number_format = MONEY
                    g = ws.cell(row=new_r, column=7)
                    size = ws.cell(row=new_r, column=6).value
                    if isinstance(size, (int, float)) and size:
                        g.value = f'=B{new_r}/F{new_r}'
                    g.number_format = MONEY
                    f['new_row'] = new_r

            elif item[0] == 'month':
                info = item[1]
                text = f"{MONTHS[info['midx']-1]} ({info['count']})"
                src = tsnap[tmonth]['cells'] if tmonth else None
                for c in range(1, max_col + 1):
                    cell = ws.cell(row=new_r, column=c)
                    cell.value = text if c == 1 else None
                    cell.hyperlink = None
                    if src is not None:
                        cell._style = copy.copy(src[c]['style'])
                        cell.number_format = src[c]['fmt']
                    if c <= 7:
                        base = cell.font
                        cell.font = Font(name='Arial', sz=19, bold=True,
                                         color=base.color if (base.color and
                                         base.color.rgb not in (None, 'FF4285F4'))
                                         else 'FF000000')
                        cell.fill = PatternFill(fill_type='solid', fgColor=GREEN)
                        if c == 1:
                            cell.alignment = Alignment(horizontal='center',
                                                       vertical=cell.alignment.vertical)
                h = tsnap[tmonth]['h'] if tmonth and tsnap[tmonth]['h'] else 24.0
                ws.row_dimensions[new_r].height = h
                new_month_rows.append((new_r, text))

            else:  # new property row
                rec = item[1]
                tp = rec.get('_tmpl') or tprop
                if tp is not None and tp not in tsnap:
                    tp = tprop
                src = tsnap[tp]['cells'] if tp else None
                vals = {1: rec['addr'], 2: rec['price'], 3: rec['bed'], 4: rec['bath'],
                        5: rec['car'], 6: rec['size'],
                        7: (f'=B{new_r}/F{new_r}'
                            if rec['price'] is not None and rec['size'] else None)}
                for c in range(1, max_col + 1):
                    cell = ws.cell(row=new_r, column=c)
                    if src is not None:
                        cell._style = copy.copy(src[c]['style'])
                    cell.value = vals.get(c)
                    cell.hyperlink = None
                    if c <= 7:
                        cell.font = Font(name='Arial', sz=19, bold=False,
                                         color=(BLUE if c == 1 else 'FF000000'),
                                         underline=None)
                        cell.number_format = MONEY if c in (2, 7) else 'General'
                        # a data row must never carry the month-header fill
                        if (cell.fill is not None and cell.fill.patternType == 'solid'
                                and cell.fill.fgColor is not None
                                and cell.fill.fgColor.rgb == GREEN):
                            cell.fill = PatternFill(fill_type=None)
                    else:
                        cell.number_format = 'General'
                h = tsnap[tprop]['h'] if tprop and tsnap[tprop]['h'] else 24.0
                ws.row_dimensions[new_r].height = h
                added_rows.append((new_r, rec['addr']))

        # tidy row heights beyond the new extent
        for r in range(len(target) + 1, sh.max_row + len(target) + 2):
            if r in ws.row_dimensions and r > len(target):
                pass

        # ---- re-merge ----------------------------------------------------
        for (mr, mc, xr, xc) in old_merges:
            nmr, nxr = omap.get(mr), omap.get(xr)
            if nmr is None or nxr is None:
                continue
            ws.merge_cells(start_row=nmr, start_column=mc,
                           end_row=nxr, end_column=xc)
        for (r, text) in new_month_rows:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)

        report['verify'].append({
            'sheet': name, 'rows_added': len(added_rows),
            'month_headers_created': [t for _, t in new_month_rows],
            'month_headers_updated': [f'{m} ({c})' for _, m, c in touched_hdr_rows],
            'month_headers_countless': countless,
            'new_rows': [{'row': r, 'addr': a} for r, a in added_rows],
            'fills': [{'row': f.get('new_row'), 'addr': f['addr'], 'price': f['price']}
                      for f in plan['fills']],
            'created_month_rows': [r for r, _ in new_month_rows],
        })

    wb.save(path)


def create_sheet(wb, skey, target, report):
    """Add a missing '<Suburb> <TYPE>' sheet at the end, copying the header row
    and column widths from an existing sheet of the same type."""
    want_house = target.endswith(' HOUSE')
    donor = None
    for cand in wb.sheetnames:            # same type first, then anything
        if cand.strip().lower() in OFF_LIMITS:
            continue
        k = sheet_key(cand)
        if k.endswith(' HOUSE') != want_house:
            continue
        sh = Sheet(wb[cand])
        if sh.headers and sh.props:
            donor = (cand, sh); break
    if donor is None:
        for cand in wb.sheetnames:
            if cand.strip().lower() in OFF_LIMITS:
                continue
            sh = Sheet(wb[cand])
            if sh.headers and sh.props:
                donor = (cand, sh); break
    if donor is None:
        return None
    dname, dsh = donor
    dws = wb[dname]
    suburb = ' '.join(w.capitalize() for w in target.split()[:-1])
    kind = target.split()[-1]
    new_name = f'{suburb} {kind}'
    ws = wb.create_sheet(title=new_name)
    for col, dim in dws.column_dimensions.items():      # column widths
        if dim.width:
            ws.column_dimensions[col].width = dim.width
    for c in range(1, dsh.max_col + 1):                 # header row
        src = dws.cell(row=1, column=c)
        tgt = ws.cell(row=1, column=c)
        tgt.value = f'{suburb} - {kind}' if c == 1 else src.value
        tgt._style = copy.copy(src._style)
        tgt.number_format = src.number_format
    if dws.row_dimensions[1].height:
        ws.row_dimensions[1].height = dws.row_dimensions[1].height
    ws.sheet_format.defaultRowHeight = dws.sheet_format.defaultRowHeight
    skey[sheet_key(new_name)] = new_name
    report['created_sheets'].append({'sheet': new_name, 'styled_after': dname})
    return new_name


def main():
    apply_changes = '--apply' in sys.argv
    exports = [a for a in sys.argv[1:] if not a.startswith('--')]
    if not exports:
        raise SystemExit('no export given')
    export_path = exports[0]
    rows = read_export(export_path)

    type_map = {}
    if '--type-map' in sys.argv:
        type_map = json.load(open(sys.argv[sys.argv.index('--type-map') + 1]))

    wb = openpyxl.load_workbook(WB_PATH)
    sheets = OrderedDict()
    skey = {}
    for name in wb.sheetnames:
        if name.strip().lower() in OFF_LIMITS:
            continue
        skey[sheet_key(name)] = name

    report = {'export': export_path, 'export_rows': len(rows), 'actions': [],
              'created_sheets': [], 'unmatched_sheets': [], 'consolidated': []}

    # ---- build records
    recs = []
    for i, r in enumerate(rows):
        addr = title_addr(str(r[0]).strip()) if r[0] is not None else ''
        suburb = str(r[1]).strip() if r[1] is not None else ''
        ptype_raw = str(r[4]).strip() if r[4] is not None else ''
        is_house = ptype_raw.upper().startswith('HOUSE')
        tkey = (str(suburb).strip().upper() + '|'
                + parse_addr(str(r[0]) if r[0] is not None else '')['full'])
        if tkey in type_map:
            is_house = (type_map[tkey] == 'HOUSE')
        d = parse_date(r[12])
        size = clean_num(r[8] if is_house else r[9])
        recs.append({
            'i': i, 'addr': addr, 'suburb': suburb, 'type': 'HOUSE' if is_house else 'UNIT',
            'ptype_raw': ptype_raw, 'bed': small_num(r[5]), 'bath': small_num(r[6]),
            'car': small_num(r[7]), 'size': size, 'price': clean_price(r[11]),
            'price_raw': str(r[11]),
            'date': d, 'date_raw': str(r[12]), 'parsed': parse_addr(addr),
            'target': sheet_key(suburb + ' ' + ('HOUSE' if is_house else 'UNIT')),
        })

    # ---- consolidate duplicates within the export (same sheet + same property)
    groups = []
    for rec in recs:
        placed = False
        for g in groups:
            if g[0]['target'] == rec['target'] and same_property(g[0]['parsed'], rec['parsed']):
                g.append(rec); placed = True; break
        if not placed:
            groups.append([rec])
    consolidated = []
    for g in groups:
        if len(g) == 1:
            consolidated.append(g[0]); continue
        best = next((x for x in g if x['price'] is not None), g[0])
        merged = dict(best)
        for fld in ('bed', 'bath', 'car', 'size'):
            if merged[fld] in (None, '-'):
                for x in g:
                    if x[fld] not in (None, '-'):
                        merged[fld] = x[fld]; break
        consolidated.append(merged)
        report['consolidated'].append({
            'addr': g[0]['addr'], 'suburb': g[0]['suburb'],
            'rows': [{'date': x['date_raw'], 'price': x['price']} for x in g],
            'kept': {'date': merged['date_raw'], 'price': merged['price']}})
    report['consolidated_count'] = len(recs) - len(consolidated)

    # Second pass: the same sale filed under two spellings of the street name
    # (same suburb/type, unit, street number, sale date, price, bed, bath).
    sheet_streets = {}
    def streets_of(target):
        if target not in sheet_streets:
            nm = skey.get(target)
            ss = set()
            if nm is not None:
                w = wb[nm]
                for rr in range(2, w.max_row + 1):
                    v = w.cell(row=rr, column=1).value
                    if (v is None or 'vacancy' in str(v).lower()
                            or month_header(w.cell(row=rr, column=1), str(v))):
                        continue
                    pp = parse_addr(str(v))
                    if pp.get('street'):
                        ss.add(pp['street'])
            sheet_streets[target] = ss
        return sheet_streets[target]

    keyed = defaultdict(list)
    for rec in consolidated:
        p = rec['parsed']
        if p.get('street') is None:
            keyed[('solo', rec['i'])].append(rec); continue
        keyed[(rec['target'], p['unit'], p['lo'], p['hi'], rec['date_raw'],
               rec['price'], str(rec['bed']), str(rec['bath']))].append(rec)
    final = []
    for k, g in keyed.items():
        if len(g) == 1 or k[0] == 'solo':
            final.extend(g); continue
        if len({r['parsed']['street'] for r in g}) == 1:
            final.extend(g); continue
        known = streets_of(g[0]['target'])
        keep = next((r for r in g if r['parsed']['street'] in known), g[0])
        merged = dict(keep)
        for fld in ('bed', 'bath', 'car', 'size', 'price'):
            if merged[fld] in (None, '-'):
                for x in g:
                    if x[fld] not in (None, '-'):
                        merged[fld] = x[fld]; break
        final.append(merged)
        report['consolidated'].append({
            'addr': keep['addr'], 'suburb': keep['suburb'],
            'note': 'same sale filed under two street names',
            'rows': [{'addr': x['addr'], 'date': x['date_raw'], 'price': x['price']} for x in g],
            'kept': {'addr': merged['addr'], 'date': merged['date_raw'],
                     'price': merged['price']}})
    final.sort(key=lambda r: r['i'])
    report['consolidated_count'] += len(consolidated) - len(final)
    consolidated = final

    # ---- plan per sheet
    plans = defaultdict(lambda: {'new': [], 'fills': []})
    model = {}
    already = 0
    for rec in consolidated:
        name = skey.get(rec['target'])
        if name is None:
            name = create_sheet(wb, skey, rec['target'], report)
            if name is None:
                report['unmatched_sheets'].append(rec['target'])
                continue
        if name not in model:
            model[name] = Sheet(wb[name])
            if model[name].tmpl_month is None or model[name].tmpl_prop is None:
                for d in report['created_sheets']:
                    if d['sheet'] == name:
                        model[name].donor = d['styled_after']
        sh = model[name]
        hit = sh.find(rec['parsed'])
        # A property RP Data has reclassified (House <-> Unit) must not be added
        # a second time to the suburb's other sheet.
        sib_name = skey.get(sibling_key(rec['target']))
        if hit is None and sib_name:
            if sib_name not in model:
                model[sib_name] = Sheet(wb[sib_name])
            sib_hit = model[sib_name].find(rec['parsed'])
            if sib_hit is not None:
                sh, name, hit = model[sib_name], sib_name, sib_hit
                report['actions'].append({
                    'kind': 'cross_sheet_match', 'sheet': sib_name, 'row': sib_hit,
                    'addr': rec['addr'],
                    'note': f"already filed on {sib_name}; export now says "
                            f"{rec['ptype_raw']!r} - not duplicated"})
        pend = None
        if hit is None:
            for q in plans[name]['new']:
                if same_property(q['parsed'], rec['parsed']):
                    pend = q; break
        if hit is not None:
            ws = sh.ws
            cur = ws.cell(row=hit, column=2).value
            blank = cur is None or str(cur).strip() == ''
            if blank and rec['price'] is not None:
                plans[name]['fills'].append({'row': hit, 'price': rec['price'],
                                             'addr': str(ws.cell(row=hit, column=1).value)})
                report['actions'].append({'kind': 'fill_price', 'sheet': name, 'row': hit,
                                          'addr': str(ws.cell(row=hit, column=1).value),
                                          'price': rec['price'], 'export_addr': rec['addr']})
            else:
                already += 1
                report['actions'].append({'kind': 'skip_existing', 'sheet': name, 'row': hit,
                                          'addr': str(ws.cell(row=hit, column=1).value),
                                          'export_addr': rec['addr'],
                                          'reason': 'price already set' if not blank
                                          else 'no usable price in export'})
        elif pend is not None:
            already += 1
            report['actions'].append({'kind': 'skip_dup_in_run', 'sheet': name,
                                      'addr': rec['addr']})
        else:
            if rec['date'] is None:
                report['actions'].append({'kind': 'skipped_no_date', 'sheet': name,
                                          'addr': rec['addr'], 'date_raw': rec['date_raw']})
                continue
            plans[name]['new'].append(rec)
            report['actions'].append({'kind': 'new_row', 'sheet': name, 'addr': rec['addr'],
                                      'month': MONTHS[rec['date'][1]-1], 'price': rec['price'],
                                      'size': rec['size'], 'bed': rec['bed'],
                                      'bath': rec['bath'], 'car': rec['car']})
    report['placeholder_prices'] = [
        {'addr': r['addr'], 'suburb': r['suburb'], 'raw': r['price_raw']}
        for r in consolidated if r['price'] is None
        and re.match(r'^[\d.]+$', r['price_raw']) and float(r['price_raw']) <= 1]
    report['already_present'] = already
    report['new_total'] = sum(len(p['new']) for p in plans.values())
    report['fills_total'] = sum(len(p['fills']) for p in plans.values())
    report['per_sheet'] = {k: {'new': len(v['new']), 'fills': len(v['fills'])}
                           for k, v in plans.items() if v['new'] or v['fills']}
    report['months_new'] = {}
    for name, p in plans.items():
        sh = model[name]
        have = {b[1] for b in sh.blocks if b[1] is not None}
        want = {rec['date'][1] for rec in p['new']}
        miss = sorted(want - have)
        if miss:
            report['months_new'][name] = [MONTHS[m-1] for m in miss]

    rpt_path = (sys.argv[sys.argv.index('--report')+1]
                if '--report' in sys.argv else '/dev/stdout')

    if not apply_changes:
        json.dump(report, open(rpt_path, 'w'), indent=1, default=str)
        return

    # ---------------------------------------------------------------- apply
    write_changes(wb, model, plans, WB_PATH, report)
    json.dump(report, open(rpt_path, 'w'), indent=1, default=str)


if __name__ == '__main__':
    main()
