#!/usr/bin/env python3
"""Relocate a property row into its correct month block.

Rebuilds the sheet as a permutation of its existing rows, so Rate PM2 formulas
and merged month headers are remapped rather than left dangling -- the same
technique update_workbook.py uses for inserts.
"""
import sys, copy, re, json
import openpyxl
sys.path.insert(0, sys.argv[1])          # folder holding update_workbook.py
import update_workbook as U

WB = sys.argv[2]

MOVES = [
    # sheet, address, target month index, drop source header if block empties
    ('Randwick UNIT',   '2/15-17 Park Avenue', 5,  False),   # 29 May 2026 sale
    ('Point Piper UNIT', '7/6 Wentworth Street', 6, True),   # 01 Jun 2026 sale
]

wb = openpyxl.load_workbook(WB)
log = []

for sheet, addr, tgt_midx, drop_empty in MOVES:
    ws = wb[sheet]
    sh = U.Sheet(ws)
    max_col = sh.max_col

    src = None
    for r, p in sh.props:
        if str(ws.cell(row=r, column=1).value).strip().lower() == addr.lower():
            src = r
            break
    if src is None:
        raise SystemExit('row not found: %s / %s' % (sheet, addr))

    src_blk = next(b for b in sh.blocks if b[0] < src <= b[4])
    tgt_blk = next((b for b in sh.blocks if b[1] == tgt_midx), None)
    if tgt_blk is None:
        raise SystemExit('no %s block in %s' % (U.MONTHS[tgt_midx-1], sheet))

    # rows of the source block that are properties, excluding the one moving
    src_remaining = [r for r, _ in sh.props if src_blk[0] < r <= src_blk[4] and r != src]
    drop_hdr = drop_empty and not src_remaining
    anchor = tgt_blk[4]                   # insert directly after the block's end

    snap = U.snapshot(ws, sh.max_row, max_col)

    # ---- new row ordering -------------------------------------------------
    order = []
    for r in range(1, sh.max_row + 1):
        if r == src or (drop_hdr and r == src_blk[0]):
            continue
        order.append(r)
        if r == anchor:
            order.append(src)
    omap = {old: new for new, old in enumerate(order, start=1)}

    # ---- header count edits (computed before rewriting) -------------------
    edits = {}
    if tgt_blk[3] is not None:
        edits[tgt_blk[0]] = re.sub(r'\(\d+\)', '(%d)' % (tgt_blk[3] + 1),
                                   str(snap[tgt_blk[0]]['cells'][1]['v']))
    if not drop_hdr and src_blk[3] is not None:
        edits[src_blk[0]] = re.sub(r'\(\d+\)', '(%d)' % (src_blk[3] - 1),
                                   str(snap[src_blk[0]]['cells'][1]['v']))

    old_merges = [(m.min_row, m.min_col, m.max_row, m.max_col)
                  for m in list(ws.merged_cells.ranges)]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))

    # ---- rewrite every row in the new order -------------------------------
    for new_r, old_r in enumerate(order, start=1):
        s = snap[old_r]
        for c in range(1, max_col + 1):
            cell = ws.cell(row=new_r, column=c)
            cd = s['cells'][c]
            v = cd['v']
            if isinstance(v, str) and v.startswith('='):
                v = U.remap_formula(v, omap)
            if c == 1 and old_r in edits:
                v = edits[old_r]
            cell.value = v
            cell._style = copy.copy(cd['style'])
            cell.number_format = cd['fmt']
            cell.hyperlink = copy.copy(cd['link']) if cd['link'] else None
        if s['h'] is not None:
            ws.row_dimensions[new_r].height = s['h']

    # blank out any rows the sheet shed
    for r in range(len(order) + 1, sh.max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.value = None
            cell.hyperlink = None

    # ---- restore merges at remapped positions -----------------------------
    for mnr, mnc, mxr, mxc in old_merges:
        if mnr != mxr:
            continue
        if mnr not in omap:
            continue
        nr = omap[mnr]
        ws.merge_cells(start_row=nr, start_column=mnc, end_row=nr, end_column=mxc)

    log.append({'sheet': sheet, 'addr': addr,
                'from_row': src, 'from_block': src_blk[2],
                'to_block': U.MONTHS[tgt_midx-1], 'to_row': omap[src],
                'header_edits': {snap[k]['cells'][1]['v']: v for k, v in edits.items()},
                'dropped_empty_header': drop_hdr and src_blk[2] or None})

wb.save(WB)
print(json.dumps(log, indent=1))
