#!/usr/bin/env python3
"""Rebuild the "Summary" tab at the end of the workbook.

Run as part of update_workbook.py, or standalone:
    python3 build_summary.py                     # rebuild from the workbook
    python3 build_summary.py --report report.json  # include "this run" figures

Everything here is derived from the suburb sheets themselves, so the tab can be
rebuilt at any time and always reflects what is actually in the workbook.

Counts include every recorded sale. Price figures can only ever cover sales
whose price RP Data has disclosed, which for recent months is well under half
of them -- every price line says how many sales it is based on.
"""
import re, sys, json, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

import update_workbook as U

SHEET = 'Summary'
OFF_LIMITS = U.OFF_LIMITS | {SHEET.lower()}
GREEN = U.GREEN
MONEY = U.MONEY
FONT = 'Arial'
SZ_TITLE, SZ_HEAD, SZ_BODY = 20, 14, 12
WIDTHS = [34, 13, 13, 13, 11, 11, 11, 30, 16]


def suburb_of(name):
    return re.sub(r'\s*-?\s*(HOUSE|UNIT)\s*$', '', name, flags=re.I).strip()


def type_of(name):
    return 'UNIT' if re.search(r'UNIT\s*$', name, flags=re.I) else 'HOUSE'


def collect(wb):
    """Every property row in the workbook, with suburb / type / month / price."""
    out = []
    for name in wb.sheetnames:
        if name.strip().lower() in OFF_LIMITS:
            continue
        ws = wb[name]
        sh = U.Sheet(ws)
        sub, typ = suburb_of(name), type_of(name)
        for r, _ in sh.props:
            price = ws.cell(row=r, column=2).value
            out.append({
                'sheet': name, 'suburb': sub, 'type': typ,
                'month': sh.month_of(r),
                'addr': ws.cell(row=r, column=1).value,
                'price': price if isinstance(price, (int, float)) else None,
            })
    return out


class Writer:
    def __init__(self, ws):
        self.ws = ws
        self.r = 1

    def _set(self, col, val, *, bold=False, size=SZ_BODY, fmt=None, align=None):
        c = self.ws.cell(row=self.r, column=col, value=val)
        c.font = Font(name=FONT, size=size, bold=bold)
        if fmt:
            c.number_format = fmt
        if align:
            c.alignment = Alignment(horizontal=align)
        return c

    def blank(self, n=1):
        self.r += n

    def title(self, text, sub=None):
        self._set(1, text, bold=True, size=SZ_TITLE)
        self.r += 1
        if sub:
            self.note(sub)

    def section(self, text, span=len(WIDTHS)):
        c = self._set(1, text, bold=True, size=SZ_HEAD, align='center')
        c.fill = PatternFill(fill_type='solid', fgColor=GREEN)
        for col in range(2, span + 1):
            cc = self.ws.cell(row=self.r, column=col)
            cc.fill = PatternFill(fill_type='solid', fgColor=GREEN)
            cc.font = Font(name=FONT, size=SZ_HEAD, bold=True)
        self.ws.merge_cells(start_row=self.r, start_column=1,
                            end_row=self.r, end_column=span)
        self.r += 1

    def row(self, values, *, bold=False, fmts=None, align=None):
        for i, v in enumerate(values, start=1):
            if v is None:
                continue
            fmt = (fmts or {}).get(i)
            self._set(i, v, bold=bold, fmt=fmt,
                      align=(align if i > 1 else None))
        self.r += 1

    def note(self, text):
        c = self._set(1, text)
        c.font = Font(name=FONT, size=SZ_BODY, italic=True, color='FF666666')
        self.r += 1

    def kv(self, label, value, fmt=None, note=None):
        self._set(1, label)
        self._set(2, value, bold=True, fmt=fmt)
        if note:
            c = self._set(4, note)
            c.font = Font(name=FONT, size=SZ_BODY, italic=True, color='FF666666')
        self.r += 1


def pct(now, before):
    """Percent change, or None when there is no sensible baseline."""
    if not before:
        return None
    return (now - before) / before


def build(wb, report=None, today=None):
    rows = collect(wb)
    months = sorted({r['month'] for r in rows if r['month']})
    if not months:
        return None
    latest, prev = months[-1], (months[-2] if len(months) > 1 else None)
    mname = lambda m: U.MONTHS[m - 1] if m else '-'

    subs = sorted({r['suburb'] for r in rows})
    this_m = [r for r in rows if r['month'] == latest]
    last_m = [r for r in rows if r['month'] == prev] if prev else []
    priced_this = [r for r in this_m if r['price']]

    if SHEET in wb.sheetnames:
        wb.remove(wb[SHEET])
    ws = wb.create_sheet(SHEET)            # always lands at the end
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    w = Writer(ws)
    stamp = (today or datetime.date.today()).strftime('%d %b %Y')
    w.title('EASTERN SUBURBS - ACTIVITY SUMMARY',
            'Rebuilt %s. %s is still filling: later exports keep adding sales to '
            'recent months, so its counts will rise.' % (stamp, mname(latest)))
    w.blank()

    # ---------------------------------------------------------------- overview
    w.section('THIS MONTH AT A GLANCE')
    w.kv('Total sales recorded (year to date)', len(rows), fmt='#,##0')
    w.kv('Sales in %s' % mname(latest), len(this_m), fmt='#,##0')
    if prev:
        w.kv('Sales in %s' % mname(prev), len(last_m), fmt='#,##0')
        d = len(this_m) - len(last_m)
        p = pct(len(this_m), len(last_m))
        w.kv('Change', d, fmt='+#,##0;-#,##0',
             note=('%+.0f%%' % (p * 100)) if p is not None else None)
    w.blank()
    houses = sum(1 for r in this_m if r['type'] == 'HOUSE')
    w.kv('Houses sold in %s' % mname(latest), houses, fmt='#,##0')
    w.kv('Units sold in %s' % mname(latest), len(this_m) - houses, fmt='#,##0')
    w.blank()
    cnt_this = {s: sum(1 for r in this_m if r['suburb'] == s) for s in subs}
    active = {s: n for s, n in cnt_this.items() if n}
    if active:
        top = max(active, key=lambda s: active[s])
        low = min(active, key=lambda s: active[s])
        w.kv('Busiest suburb', '%s (%d sales)' % (top, active[top]))
        w.kv('Quietest suburb with sales', '%s (%d)' % (low, active[low]))
    quiet = [s for s in subs if not cnt_this[s]]
    w.kv('Suburbs with no sales', ', '.join(quiet) if quiet else 'none')
    w.blank()
    if priced_this:
        hi = max(priced_this, key=lambda r: r['price'])
        lo = min(priced_this, key=lambda r: r['price'])
        for label, rec in (('Highest sale', hi), ('Lowest sale', lo)):
            w.kv(label, rec['price'], fmt=MONEY,
                 note='%s, %s (%s)' % (rec['addr'], rec['suburb'],
                                       rec['type'].lower()))
        w.note('Highest, lowest and the ranges below cover only the %d of %d %s '
               'sales whose price has been disclosed.'
               % (len(priced_this), len(this_m), mname(latest)))
    w.blank()

    # ---------------------------------------------------------------- suburbs
    w.section('BY SUBURB')
    w.row(['Suburb', mname(latest), mname(prev) if prev else '', 'Change',
           'YTD', 'Share', 'Avg/mo', '%s price range' % mname(latest), ''],
          bold=True, align='center')
    for s in subs:
        ytd = sum(1 for r in rows if r['suburb'] == s)
        n_now, n_bef = cnt_this[s], sum(1 for r in last_m if r['suburb'] == s)
        p = [r['price'] for r in this_m if r['suburb'] == s and r['price']]
        rng = ('%s - %s' % (f'${min(p):,.0f}', f'${max(p):,.0f}')
               if len(p) > 1 else (f'${p[0]:,.0f}' if p else '-'))
        note = ('%d priced' % len(p)) if p else 'no prices yet'
        w.row([s, n_now, n_bef if prev else None, n_now - n_bef,
               ytd, ytd / len(rows), ytd / latest, rng, note],
              fmts={4: '+#,##0;-#,##0;0', 6: '0.0%', 7: '0.0'}, align='center')
    w.row(['TOTAL', len(this_m), len(last_m) if prev else None,
           len(this_m) - len(last_m), len(rows), 1.0, len(rows) / latest, '', ''],
          bold=True, fmts={4: '+#,##0;-#,##0;0', 6: '0.0%', 7: '0.0'}, align='center')
    w.blank()

    # ---------------------------------------------------------------- months
    w.section('SALES BY MONTH')
    w.row(['Month', 'Sales', 'Change', '', '', '', '', '', ''],
          bold=True, align='center')
    for i, m in enumerate(months):
        n = sum(1 for r in rows if r['month'] == m)
        before = sum(1 for r in rows if r['month'] == months[i - 1]) if i else None
        w.row([mname(m) + (' (partial)' if m == latest else ''), n,
               (n - before) if before is not None else None],
              fmts={3: '+#,##0;-#,##0;0'}, align='center')
    w.blank()

    # ---------------------------------------------------------------- run log
    if report:
        w.section('LAST UPDATE')
        w.kv('Export processed', str(report.get('export', '-')))
        w.kv('Rows in export', report.get('export_rows', 0), fmt='#,##0')
        w.kv('Already recorded', report.get('already_present', 0), fmt='#,##0')
        w.kv('New rows added', report.get('new_total', 0), fmt='#,##0')
        w.kv('Late prices filled in', report.get('fills_total', 0), fmt='#,##0')
        w.kv('Flagged for review', report.get('second_sales_total', 0), fmt='#,##0')
    return ws


def refresh(path=None, report=None):
    """Rebuild the Summary tab in the saved workbook. Safe to call repeatedly."""
    path = path or U.WB_PATH
    wb = openpyxl.load_workbook(path)
    build(wb, report)
    wb.save(path)
    return path


def main():
    rpt = None
    if '--report' in sys.argv:
        rpt = json.load(open(sys.argv[sys.argv.index('--report') + 1]))
    print('Summary rebuilt in %s' % refresh(U.WB_PATH, rpt))


if __name__ == '__main__':
    main()
